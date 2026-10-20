import json, pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.comments import Comment
RED = PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE")
def run_rules(df, rules):
    errs = []
    for r in rules:
        t, nm = r["type"], r["name"]
        cols = [c.strip() for c in json.loads(r["params"] or "{}").get("columns", "").split(",") if c.strip()]
        if t == "required_columns":
            errs += [(None, c, nm, f"Missing required column '{c}'") for c in cols if c not in df.columns]
            continue
        use = [c for c in (cols or list(df.columns)) if c in df.columns]
        if t == "null_check":
            for c in use: errs += [(i, c, nm, "Null / empty value") for i in df.index[df[c].isna()]]
        elif t == "duplicate_check" and use:
            errs += [(i, use[0], nm, "Duplicate row") for i in df.index[df.duplicated(subset=use, keep="first")]]
        elif t == "outlier_check":
            for c in use:
                s = pd.to_numeric(df[c], errors="coerce")
                if s.notna().sum() < 4: continue
                q1, q3 = s.quantile(.25), s.quantile(.75); k = (q3 - q1) * 1.5
                errs += [(i, c, nm, f"Outlier ({s[i]})") for i in s.index[(s < q1 - k) | (s > q3 + k)]]
        elif t == "whitespace_check":
            for c in use:
                m = df[c].apply(lambda v: isinstance(v, str) and v != v.strip())
                errs += [(i, c, nm, "Leading/trailing whitespace") for i in df.index[m]]
        elif t == "negative_check":
            for c in use:
                s = pd.to_numeric(df[c], errors="coerce")
                errs += [(i, c, nm, f"Negative value ({s[i]})") for i in s.index[s < 0]]
    return errs
def annotate(src, out, df, errs):
    wb = load_workbook(src); ws = wb.active; cols = list(df.columns)
    for i, c, nm, msg in errs:
        if i is None or c not in cols: continue
        cell = ws.cell(row=int(i) + 2, column=cols.index(c) + 1)
        cell.fill = RED
        cell.comment = Comment(f"{nm}: {msg}", "DQM") if not cell.comment else Comment(cell.comment.text + f"\n{nm}: {msg}", "DQM")
    es = wb.create_sheet("Errors"); es.append(["Row", "Column", "Rule", "Issue"])
    for i, c, nm, msg in errs: es.append(["-" if i is None else int(i) + 2, c, nm, msg])
    wb.save(out)
